from __future__ import annotations

import io
import unittest

from docx import Document

from src.resume_parser import (
    clean_extracted_text,
    extract_resume_text,
    get_resume_statistics,
)


class TestResumeParser(unittest.TestCase):

    def test_txt_extraction(self):
        content = b"""
        John Doe
        Python Developer
        Skills: Python, SQL, Machine Learning
        """

        file_obj = io.BytesIO(content)

        result = extract_resume_text(
            file_obj,
            "resume.txt",
        )

        self.assertIn("John Doe", result)
        self.assertIn("Python Developer", result)
        self.assertIn("Machine Learning", result)

    def test_docx_extraction(self):
        document = Document()

        document.add_paragraph("John Doe")
        document.add_paragraph("AI/ML Engineer")
        document.add_paragraph("Skills: Python, TensorFlow")

        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)

        result = extract_resume_text(
            buffer,
            "resume.docx",
        )

        self.assertIn("John Doe", result)
        self.assertIn("AI/ML Engineer", result)
        self.assertIn("TensorFlow", result)

    def test_text_cleaning(self):
        text = """
        John     Doe


        Python       Developer
        """

        result = clean_extracted_text(text)

        self.assertEqual(
            result,
            "John Doe\nPython Developer",
        )

    def test_resume_statistics(self):
        text = """
        John Doe
        Python Developer
        Machine Learning Engineer
        """

        statistics = get_resume_statistics(text)

        self.assertGreater(
            statistics["words"],
            0,
        )

        self.assertGreater(
            statistics["characters"],
            0,
        )

        self.assertEqual(
            statistics["lines"],
            3,
        )

    def test_unsupported_file(self):
        file_obj = io.BytesIO(
            b"Some resume content"
        )

        with self.assertRaises(ValueError):
            extract_resume_text(
                file_obj,
                "resume.jpg",
            )


if __name__ == "__main__":
    unittest.main()